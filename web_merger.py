import streamlit as st
from PyPDF2 import PdfMerger, PdfReader
from docx import Document
import io

# 1. Page Configuration
st.set_page_config(page_title="PDF Merger Pro", page_icon="📄", layout="centered")

# 2. Inject Custom CSS for Professional UI
st.markdown("""
    <style>
        /* Custom Logo Banner */
        .logo-container {
            display: flex;
            align-items: center;
            justify-content: center;
            padding: 25px;
            background: linear-gradient(135deg, #0f2027 0%, #203a43 50%, #2c5364 100%);
            border-radius: 12px;
            margin-bottom: 30px;
            color: white;
            box-shadow: 0 4px 15px rgba(0,0,0,0.2);
        }
        .logo-icon { font-size: 40px; margin-right: 15px; }
        .logo-text { font-size: 34px; font-weight: 800; letter-spacing: 1px; margin: 0; }
        .logo-pro { color: #00d2ff; }
        
        /* Attractive Upload Zone */
        [data-testid="stFileUploadDropzone"] {
            border: 2px dashed #00d2ff !important;
            border-radius: 15px !important;
            background-color: rgba(0, 210, 255, 0.05) !important;
            padding: 30px !important;
            transition: all 0.3s ease-in-out;
        }
        [data-testid="stFileUploadDropzone"]:hover {
            background-color: rgba(0, 210, 255, 0.15) !important;
            border-color: #2c5364 !important;
            transform: scale(1.01);
        }
        
        /* File Numbering Badge */
        .file-header {
            display: flex;
            align-items: center;
            background-color: #f1f3f6;
            padding: 10px 15px;
            border-radius: 8px;
            margin-bottom: 10px;
            color: #1e1e1e;
            font-weight: 600;
        }
        .file-badge {
            background: linear-gradient(135deg, #203a43, #2c5364);
            color: white;
            width: 30px;
            height: 30px;
            display: flex;
            align-items: center;
            justify-content: center;
            border-radius: 50%;
            font-size: 16px;
            font-weight: bold;
            margin-right: 15px;
        }
        
        /* Dark mode compatibility for file header */
        @media (prefers-color-scheme: dark) {
            .file-header { background-color: #1e1e1e; color: #ffffff; border: 1px solid #333; }
        }
    </style>
""", unsafe_allow_html=True)

# 3. Custom Logo Render
st.markdown("""
    <div class="logo-container">
        <div class="logo-icon">📄</div>
        <div class="logo-text">PDF Merger <span class="logo-pro">Pro</span></div>
    </div>
""", unsafe_allow_html=True)

st.write("Upload your files below, configure the pages you want to keep, and merge them instantly.")

# 4. File Upload Section
uploaded_files = st.file_uploader("Drop your PDF files here", type="pdf", accept_multiple_files=True)

if uploaded_files:
    st.markdown("### ⚙️ Configure PDF Sequence")
    
    merge_instructions = []
    
    for i, file in enumerate(uploaded_files, 1):
        file.seek(0)
        pdf_reader = PdfReader(file)
        total_pages = len(pdf_reader.pages)
        
        # Custom HTML Header for each file with a number badge
        st.markdown(f"""
            <div class="file-header">
                <div class="file-badge">{i}</div>
                <div>{file.name} <span style="font-weight:normal; opacity:0.8; font-size:14px;">({total_pages} Pages)</span></div>
            </div>
        """, unsafe_allow_html=True)
        
        # Settings container for this specific file
        with st.container():
            col1, col2 = st.columns(2)
            
            with col1:
                pages_to_keep = st.number_input(
                    "Pages to extract:",
                    min_value=1,
                    max_value=total_pages,
                    value=total_pages,
                    key=f"count_{i}",
                    help="Choose how many pages you want to take from this document."
                )
            
            with col2:
                section = st.selectbox(
                    "Extract from:",
                    options=["Top (Beginning)", "Middle", "Bottom (End)"],
                    key=f"pos_{i}",
                    disabled=(pages_to_keep == total_pages)
                )
            
            # Math logic for page slicing
            if pages_to_keep == total_pages:
                start_page, end_page = 0, total_pages
            elif section == "Top (Beginning)":
                start_page, end_page = 0, pages_to_keep
            elif section == "Bottom (End)":
                start_page, end_page = total_pages - pages_to_keep, total_pages
            else:
                start_page = (total_pages - pages_to_keep) // 2
                end_page = start_page + pages_to_keep
                
            st.caption(f"🔹 *Including pages **{start_page + 1} to {end_page}** of this document.*")
            st.write("") # Spacer
            
            merge_instructions.append({"file": file, "start": start_page, "end": end_page})
            
    # 5. Big Merge Button
    if len(uploaded_files) > 1:
        st.markdown("<br>", unsafe_allow_html=True)
        if st.button("✨ Merge PDFs Now", type="primary", use_container_width=True):
            with st.spinner("Stitching PDFs together..."):
                merger = PdfMerger()
                for instruction in merge_instructions:
                    instruction["file"].seek(0)
                    merger.append(instruction["file"], pages=(instruction["start"], instruction["end"]))
                
                merged_pdf = io.BytesIO()
                merger.write(merged_pdf)
                merger.close()
                merged_pdf.seek(0)
                
                st.session_state['merged_pdf'] = merged_pdf
                st.success("🎉 Documents merged successfully!")

# 6. Stylish Download Options
if 'merged_pdf' in st.session_state:
    st.markdown("---")
    st.markdown("### 📥 Download Your File")
    col1, col2 = st.columns(2)
    
    with col1:
        st.download_button(
            label="📄 Download as PDF",
            data=st.session_state['merged_pdf'],
            file_name="merged_document.pdf",
            mime="application/pdf",
            use_container_width=True
        )
        
    with col2:
        # Generate Word doc on the fly
        doc = Document()
        doc.add_heading("Merged PDF Text", level=1)
        st.session_state['merged_pdf'].seek(0)
        pdf_reader = PdfReader(st.session_state['merged_pdf'])
        
        for page_num, page in enumerate(pdf_reader.pages, 1):
            text = page.extract_text()
            if text:
                doc.add_paragraph(f"--- Page {page_num} ---")
                doc.add_paragraph(text)
        
        word_buffer = io.BytesIO()
        doc.save(word_buffer)
        word_buffer.seek(0)
        
        st.download_button(
            label="📝 Download as Word Document",
            data=word_buffer,
            file_name="merged_document.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )